import os
import re
from docx import Document
from PyPDF2 import PdfReader
from typing import Dict, List, Any

class DocumentProcessor:
    def __init__(self):
        self.supported_types = {
            'application/pdf': self._process_pdf,
            'application/vnd.openxmlformats-officedocument.wordprocessingml.document': self._process_docx,
            'text/plain': self._process_txt
        }

    async def process_document(self, file_path: str, mime_type: str) -> Dict[str, Any]:
        """
        Processa um documento e retorna seus metadados e conteúdo estruturado.
        """
        if mime_type not in self.supported_types:
            raise ValueError(f"Tipo de arquivo não suportado: {mime_type}")

        processor = self.supported_types[mime_type]
        return await processor(file_path)

    async def _process_pdf(self, file_path: str) -> Dict[str, Any]:
        """
        Processa um arquivo PDF.
        """
        chapters = []
        metadata = {}
        current_chapter = {"title": "Capítulo 1", "paragraphs": []}
        
        try:
            with open(file_path, 'rb') as file:
                pdf = PdfReader(file)
                metadata = {
                    "num_pages": len(pdf.pages),
                    "title": pdf.metadata.get('/Title', ''),
                    "author": pdf.metadata.get('/Author', ''),
                    "subject": pdf.metadata.get('/Subject', '')
                }

                for page_num, page in enumerate(pdf.pages):
                    text = page.extract_text()
                    paragraphs = [p.strip() for p in text.split('\n\n') if p.strip()]
                    
                    for paragraph in paragraphs:
                        # Detectar possíveis títulos de capítulo
                        if re.match(r'^(Chapter|Capítulo)\s+\d+', paragraph, re.IGNORECASE):
                            if current_chapter["paragraphs"]:
                                chapters.append(current_chapter)
                            current_chapter = {"title": paragraph, "paragraphs": []}
                        else:
                            current_chapter["paragraphs"].append(paragraph)

                if current_chapter["paragraphs"]:
                    chapters.append(current_chapter)

        except Exception as e:
            raise Exception(f"Erro ao processar PDF: {str(e)}")

        return {
            "metadata": metadata,
            "chapters": chapters
        }

    async def _process_docx(self, file_path: str) -> Dict[str, Any]:
        """
        Processa um arquivo DOCX.
        """
        chapters = []
        metadata = {}
        current_chapter = {"title": "Capítulo 1", "paragraphs": []}

        try:
            doc = Document(file_path)
            metadata = {
                "title": doc.core_properties.title or '',
                "author": doc.core_properties.author or '',
                "subject": doc.core_properties.subject or '',
                "num_paragraphs": len(doc.paragraphs)
            }

            for paragraph in doc.paragraphs:
                text = paragraph.text.strip()
                if not text:
                    continue

                # Detectar possíveis títulos de capítulo
                if paragraph.style.name.startswith('Heading') or \
                   re.match(r'^(Chapter|Capítulo)\s+\d+', text, re.IGNORECASE):
                    if current_chapter["paragraphs"]:
                        chapters.append(current_chapter)
                    current_chapter = {"title": text, "paragraphs": []}
                else:
                    current_chapter["paragraphs"].append(text)

            if current_chapter["paragraphs"]:
                chapters.append(current_chapter)

        except Exception as e:
            raise Exception(f"Erro ao processar DOCX: {str(e)}")

        return {
            "metadata": metadata,
            "chapters": chapters
        }

    async def _process_txt(self, file_path: str) -> Dict[str, Any]:
        """
        Processa um arquivo TXT.
        """
        chapters = []
        metadata = {
            "encoding": "utf-8",
            "size": os.path.getsize(file_path)
        }
        current_chapter = {"title": "Capítulo 1", "paragraphs": []}

        try:
            with open(file_path, 'r', encoding='utf-8') as file:
                content = file.read()
                paragraphs = [p.strip() for p in content.split('\n\n') if p.strip()]

                for paragraph in paragraphs:
                    # Detectar possíveis títulos de capítulo
                    if re.match(r'^(Chapter|Capítulo)\s+\d+', paragraph, re.IGNORECASE):
                        if current_chapter["paragraphs"]:
                            chapters.append(current_chapter)
                        current_chapter = {"title": paragraph, "paragraphs": []}
                    else:
                        current_chapter["paragraphs"].append(paragraph)

                if current_chapter["paragraphs"]:
                    chapters.append(current_chapter)

        except Exception as e:
            raise Exception(f"Erro ao processar TXT: {str(e)}")

        return {
            "metadata": metadata,
            "chapters": chapters
        }
